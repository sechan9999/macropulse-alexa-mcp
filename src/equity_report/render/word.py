"""Word (.docx) research note -> bytes.

Style rules adapted from stock-report-skills' word-report-style: navy table headers with white bold
text, zebra rows, centred cells, rows never split across pages, headings kept with the next block,
6.5-inch charts, 1.5 cm / 2 cm margins, body indent 0.5 cm, ASCII file names. Signed numbers use US
colours (green up / red down). Charts are inserted under the matching narrative heading so no chart
appears without its explanation.
"""
from __future__ import annotations

import io
import re

import numpy as np
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls, qn
from docx.shared import Cm, Inches, Pt, RGBColor

from ..engine.analysis import Report
from ..engine.dcf import FOMC_NOTE

FONT, NAVY = "Calibri", "1a1a4e"
GREEN, RED, GREY = RGBColor(0x16, 0x80, 0x3c), RGBColor(0xC0, 0x1C, 0x1C), RGBColor(0x59, 0x59, 0x59)


class _Doc:
    def __init__(self):
        self.doc = Document()
        st = self.doc.styles["Normal"]
        st.font.name, st.font.size = FONT, Pt(10.5)
        st.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
        for i in (1, 2, 3):
            hs = self.doc.styles[f"Heading {i}"]
            hs.font.name = FONT
            hs.font.color.rgb = RGBColor(0x1A, 0x1A, 0x4E)
        for s in self.doc.sections:
            s.top_margin = s.bottom_margin = Cm(1.5)
            s.left_margin = s.right_margin = Cm(2)

    def heading(self, text, level=1):
        h = self.doc.add_heading(text, level=level)
        h.paragraph_format.keep_with_next = True

    def para(self, text, indent=True, size=None, color=None, bold=False):
        p = self.doc.add_paragraph()
        if indent:
            p.paragraph_format.left_indent = Cm(0.5)
        for part in re.split(r"(\*\*[^*]+\*\*)", text):
            if not part:
                continue
            is_b = part.startswith("**") and part.endswith("**")
            run = p.add_run(part[2:-2] if is_b else part)
            run.bold = is_b or bold
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color
        return p

    def image(self, png: bytes | None, caption: str | None = None):
        if not png:
            return
        self.doc.add_picture(io.BytesIO(png), width=Inches(6.5))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            self.para(caption, indent=False, size=8.5, color=GREY).alignment = WD_ALIGN_PARAGRAPH.CENTER

    def table(self, header, rows, signed_cols=()):
        t = self.doc.add_table(rows=1 + len(rows), cols=len(header))
        t.alignment, t.style = WD_TABLE_ALIGNMENT.CENTER, "Table Grid"
        for j, h in enumerate(header):
            cell = t.rows[0].cells[j]
            cell.text = ""
            r = cell.paragraphs[0].add_run(str(h))
            r.bold, r.font.size, r.font.color.rgb = True, Pt(10), RGBColor(0xFF, 0xFF, 0xFF)
            cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{NAVY}"/>'))
        for i, row in enumerate(rows, start=1):
            for j, v in enumerate(row):
                cell = t.rows[i].cells[j]
                cell.text = ""
                r = cell.paragraphs[0].add_run(str(v))
                r.font.size = Pt(9.5)
                if j in signed_cols and isinstance(v, str):
                    if v.strip().startswith("+"):
                        r.font.color.rgb = GREEN
                    elif v.strip().startswith(("-", "−")):
                        r.font.color.rgb = RED
                if i % 2 == 0:
                    cell._tc.get_or_add_tcPr().append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="f2f2f2"/>'))
        for row in t.rows:
            row._tr.get_or_add_trPr().append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.keep_with_next = True
        self.doc.add_paragraph()

    def md_table(self, lines):
        rows = [[c.strip() for c in ln.strip().strip("|").split("|")] for ln in lines
                if not re.match(r"^\|?\s*:?-{3,}", ln.strip())]
        if len(rows) >= 2:
            self.table(rows[0], rows[1:], signed_cols=range(len(rows[0])))

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def _pct(x):
    return "n/a" if x is None or (isinstance(x, float) and np.isnan(x)) else f"{x * 100:+.1f}%"


def build(rep: Report, charts: dict[str, bytes]) -> bytes:
    D, ctx, val = _Doc(), rep.ctx, rep.val
    c, rb, v, rg = ctx["company"], ctx["rule_based"], ctx["valuation"], ctx["macro_regime"]
    lv = rb["levels"]
    D.doc.add_heading(f"{c['name']} ({c['ticker']}) — Equity Report", 0)
    D.para(f"Price date {c['as_of']} · Sources: {'; '.join(ctx['data_sources'])} · Narrative: {rep.narrative_method}",
           indent=False, size=9, color=GREY)
    if c["synthetic_data"]:
        D.para("SYNTHETIC DEMO DATA — not a real company.", indent=False, bold=True, color=RED)
    D.heading("Summary", 1)
    D.table(["Item", "Value", "vs. price"], [
        ["Share price", f"${c['price']:,.2f}", ""],
        ["Rule-based rating", f"{rb['rating']} (score {rb['score']:+.1f})", ""],
        ["DCF value (base)", f"${v['dcf_base_per_share']:,.2f}", _pct(v["dcf_base_upside"])],
        ["Probability-weighted value", f"${v['expected_per_share']:,.2f}", _pct(v["expected_upside"])],
        ["Macro regime", f"{rg['regime']} — bear/base/bull {rg['probs']['Bear']:.0%}/{rg['probs']['Base']:.0%}/"
                         f"{rg['probs']['Bull']:.0%}", ""],
        ["Stop / target 1 / target 2", f"${lv['stop']:,.2f} / ${lv['target1']:,.2f} / ${lv['target2']:,.2f}",
         f"{_pct(lv['stop'] / c['price'] - 1)} / {_pct(lv['target1'] / c['price'] - 1)} / {_pct(lv['target2'] / c['price'] - 1)}"],
        ["Pattern screen (daily)", f"{rep.screen['overall']} (bull {rep.screen['bull_score']} : bear {rep.screen['bear_score']})", ""],
    ], signed_cols=(2,))
    D.para(f"Rating basis: {rb['basis']}", size=9)

    fin = val["history"]
    inserts = {
        "1.": lambda: D.table(["USD m"] + [f"FY{y}" for y in fin.index], [
            [lab] + [("-" if np.isnan(fin.loc[y, k]) else f"{fin.loc[y, k] / 1e6:,.0f}") for y in fin.index]
            for k, lab in (("revenue", "Revenue"), ("operating_income", "Operating income"), ("net_income", "Net income"),
                           ("operating_cf", "Cash from operations"), ("capex", "CapEx"), ("fcf", "Free cash flow"))
            if k in fin]),
        "Monthly": lambda: D.image(charts.get("Monthly"), "Monthly, 3 years: candles, moving averages, support/resistance, patterns, MACD, KD"),
        "Weekly": lambda: D.image(charts.get("Weekly"), "Weekly, 12 months"),
        "Daily": lambda: D.image(charts.get("Daily"), "Daily, 6 months"),
        "4.": lambda: (
            D.table(["Scenario", "Value / share", "vs. price", f"Probability ({rg['regime']})"],
                    [[k, f"${s['per_share']:,.2f}", _pct(s["upside"]), f"{s['prob']:.0%}"] for k, s in v["scenarios"].items()],
                    signed_cols=(2,)),
            D.table(["WACC \\ g"] + [f"{g:.1%}" for g in val["sensitivity"].columns],
                    [[f"{w:.2%}"] + [("-" if np.isnan(x) else f"${x:,.0f}") for x in row]
                     for w, row in zip(val["sensitivity"].index, val["sensitivity"].values)]),
            D.table(["FOMC scenario (tab 13)", "WACC", "g", "Value / share", "vs. price", "Tab-13 price shock"],
                    [[x["title"], f"{x['wacc']:.2%}", f"{x['terminal_growth']:.1%}",
                      "-" if not np.isfinite(x["per_share"]) else f"${x['per_share']:,.2f}", _pct(x["upside"]),
                      _pct(x["tab13_price_shock"])] for x in val.get("fomc_overlay", [])],
                    signed_cols=(4, 5)),
            D.para(FOMC_NOTE, size=9, color=GREY)),
    }
    lines = (rep.narrative or "").splitlines()
    i, first = 0, True
    while i < len(lines):
        ln = lines[i].rstrip()
        if ln.startswith("## "):
            if not first:
                D.page_break()
            first = False
            title = ln[3:].strip()
            D.heading(title, 1)
            if title.split()[0] in inserts:
                inserts[title.split()[0]]()
        elif ln.startswith("### "):
            title = ln[4:].strip()
            D.heading(title, 2)
            for k in ("Monthly", "Weekly", "Daily"):
                if title.startswith(k):
                    inserts[k]()
        elif ln.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            D.md_table(block)
            continue
        elif ln.strip():
            D.para(re.sub(r"^[-*•]\s+", "", ln.strip().lstrip("> ")))
        i += 1

    D.page_break()
    D.heading("Appendix: assumptions and sources", 1)
    a = val["assumptions"]
    D.table(["Assumption", "Value"], [
        ["Risk-free rate", f"{a['risk_free_used']:.2%} ({a['risk_free_source']})"],
        ["Equity risk premium", f"{a['erp_base']:.2%} base {rg['erp_shift']:+.2%} regime shift = {a['erp_used']:.2%}"],
        ["Beta (Blume-adjusted)", f"{a['beta_used']:.2f}"], ["Pre-tax cost of debt", f"{a['cost_of_debt_used']:.2%}"],
        ["Tax rate", f"{a['tax_rate']:.0%}"], ["WACC", f"{a['wacc']:.2%}"],
        ["Year-1 growth -> terminal", f"{a['growth_high_used']:.1%} -> {a['terminal_growth']:.1%}"],
        ["FCF margin", f"{a['fcf_margin_used']:.1%}"],
    ])
    D.para("Sources: " + "; ".join(ctx["data_sources"]) + ". News, consensus estimates and the earnings calendar are not connected.",
           size=9)
    D.para("Not investment advice. Results depend heavily on the assumptions above.", size=9)
    buf = io.BytesIO()
    D.doc.save(buf)
    return buf.getvalue()
