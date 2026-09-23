"""Formula-driven Excel model (openpyxl) -> bytes.

Colour code: blue = hard-coded input, black = formula, green = link to another sheet,
yellow fill = key assumption the user is expected to change. Amounts in USD millions,
per-share values in USD. Changing any yellow cell recalculates DCF, sensitivity, scenarios
and the summary. A test recalculates this workbook in LibreOffice and checks it matches the
Python engine to 1e-6.
"""
from __future__ import annotations

import io

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter as L

from ..engine.analysis import Report
from ..engine.dcf import FOMC_NOTE

FONT = "Arial"
BLUE, GREEN, BLACK, WHITE = "0000FF", "008000", "000000", "FFFFFF"
HDR = PatternFill("solid", fgColor="1A1A4E")
KEY = PatternFill("solid", fgColor="FFFF00")
STRIPE = PatternFill("solid", fgColor="F2F2F2")
F_M = '#,##0;(#,##0);"-"'
F_PS = '$#,##0.00;($#,##0.00);"-"'
F_PCT = '0.0%;(0.0%);"-"'
F_X = '0.0"x"'
M = 1e6
FIN_ROWS = [("revenue", "Revenue"), ("operating_income", "Operating income"), ("net_income", "Net income"),
            ("eps_diluted", "Diluted EPS ($)"), ("total_assets", "Total assets"), ("total_liabilities", "Total liabilities"),
            ("total_equity", "Shareholders' equity"), ("cash", "Cash & equivalents"), ("debt", "Total debt"),
            ("operating_cf", "Cash from operations"), ("capex", "Capital expenditures"),
            ("depreciation", "Depreciation & amortization"), ("interest_expense", "Interest expense"),
            ("dividends", "Dividends paid")]


def _f(color=BLACK, bold=False, size=10):
    return Font(name=FONT, color=color, bold=bold, size=size)


def _title(ws, text, sub=None):
    ws["A1"], ws["A1"].font = text, _f(bold=True, size=14)
    if sub:
        ws["A2"], ws["A2"].font = sub, _f("595959", size=9)


def _hdr(ws, row, values, col=1):
    for j, v in enumerate(values):
        c = ws.cell(row=row, column=col + j, value=v)
        c.font, c.fill = _f(WHITE, True), HDR
        c.alignment = Alignment(horizontal="center", vertical="center")


def _put(ws, ref, value, fmt=None, color=BLACK, bold=False, fill=None, note=None):
    c = ws[ref]
    c.value, c.font = value, _f(color, bold)
    if fmt:
        c.number_format = fmt
    if fill:
        c.fill = fill
    if note:
        c.comment = Comment(note, "MacroPulse")
    return c


def _num(v):
    return None if v is None or (isinstance(v, float) and not np.isfinite(v)) or pd.isna(v) else float(v)


def build(rep: Report, chart_png: bytes | None = None) -> bytes:
    b, val, ctx = rep.bundle, rep.val, rep.ctx
    fin: pd.DataFrame = val["history"]
    years = list(fin.index.astype(str))
    ny = len(years)
    A, rg = val["assumptions"], val["regime"]
    N = int(A["projection_years"])
    last_col = L(1 + ny)
    src = "; ".join(b.get("sources", []))
    wb = Workbook()

    # ── Financials ──
    wf = wb.active
    wf.title = "Financials"
    _title(wf, f"{b['name']} ({b['ticker']}) — annual financials (USD millions)", f"Source: {src}")
    _hdr(wf, 4, ["Item (USD m)"] + [f"FY{y}" for y in years])
    wf.cell(row=5, column=1, value="Fiscal year end").font = _f()
    for j, y in enumerate(years):
        c = wf.cell(row=5, column=2 + j, value=str(fin.loc[y, "fy_end"]) if "fy_end" in fin else y)
        c.font, c.alignment = _f(BLUE), Alignment(horizontal="right")
    R = {}
    for i, (k, label) in enumerate(FIN_ROWS):
        r = 6 + i
        R[k] = r
        wf.cell(row=r, column=1, value=label).font = _f()
        for j, y in enumerate(years):
            v = _num(fin.loc[y, k]) if k in fin.columns else None
            if v is not None and k != "eps_diluted":
                v /= M
            c = wf.cell(row=r, column=2 + j, value=v)
            c.font, c.number_format = _f(BLUE), ('$0.00' if k == "eps_diluted" else F_M)
    r0 = 6 + len(FIN_ROWS) + 1
    _hdr(wf, r0, ["Ratios & derived"] + [f"FY{y}" for y in years])
    ratios = [
        ("Revenue growth", lambda c, p: f'=IFERROR({c}{R["revenue"]}/{p}{R["revenue"]}-1,"")' if p else None, F_PCT),
        ("Operating margin", lambda c, p: f'=IFERROR({c}{R["operating_income"]}/{c}{R["revenue"]},"")', F_PCT),
        ("Net margin", lambda c, p: f'=IFERROR({c}{R["net_income"]}/{c}{R["revenue"]},"")', F_PCT),
        ("ROE", lambda c, p: f'=IFERROR({c}{R["net_income"]}/{c}{R["total_equity"]},"")', F_PCT),
        ("Liabilities / equity", lambda c, p: f'=IFERROR({c}{R["total_liabilities"]}/{c}{R["total_equity"]},"")', F_X),
        ("Free cash flow (CFO - CapEx)", lambda c, p: f'={c}{R["operating_cf"]}-{c}{R["capex"]}', F_M),
        ("FCF margin", lambda c, p: f'=IFERROR(({c}{R["operating_cf"]}-{c}{R["capex"]})/{c}{R["revenue"]},"")', F_PCT),
        ("EBITDA (EBIT + D&A)", lambda c, p: f'={c}{R["operating_income"]}+{c}{R["depreciation"]}', F_M),
        ("Net debt", lambda c, p: f'={c}{R["debt"]}-{c}{R["cash"]}', F_M),
    ]
    RR = {}
    for i, (label, fn, fmt) in enumerate(ratios):
        r = r0 + 1 + i
        RR[label] = r
        wf.cell(row=r, column=1, value=label).font = _f()
        for j in range(ny):
            fml = fn(L(2 + j), L(1 + j) if j else None)
            if fml:
                c = wf.cell(row=r, column=2 + j, value=fml)
                c.font, c.number_format = _f(), fmt
    wf.column_dimensions["A"].width = 30
    for j in range(ny):
        wf.column_dimensions[L(2 + j)].width = 13
    wf.freeze_panes = "B6"

    # ── Assumptions ──
    wa = wb.create_sheet("Assumptions")
    _title(wa, "Assumptions — edit the yellow cells", "Blue = input · Green = link · Black = formula")
    _hdr(wa, 4, ["Item", "Value", "Basis / note"])
    rev_r, fcf_r = R["revenue"], RR["FCF margin"]
    first3 = L(max(2, 2 + ny - 3))
    be = A["beta_est"]
    g_cell = (f"=MIN(MAX((Financials!{last_col}{rev_r}/Financials!B{rev_r})^(1/{max(ny - 1, 1)})-1,-0.05),0.2)"
              if A.get("growth_high") is None else A["growth_high"])
    m_cell = (f"=AVERAGE(Financials!{first3}{fcf_r}:{last_col}{fcf_r})" if A.get("fcf_margin") is None else A["fcf_margin"])
    rows = [
        (5, "Company", b["name"], None, BLUE, False, f"CIK {b.get('cik') or '-'}"),
        (6, "Ticker", b["ticker"], None, BLUE, False, ""),
        (7, "Price date", b["as_of"], None, BLUE, False, "last close in the price history"),
        (8, "Share price ($)", b["last_price"], F_PS, BLUE, False, "Yahoo Finance close"),
        (9, "Shares outstanding", b["shares"], "#,##0", BLUE, False, "dei:EntityCommonStockSharesOutstanding (all classes)"),
        (10, "Market cap ($m)", "=B8*B9/1E6", F_M, BLACK, False, ""),
        (12, "Risk-free rate", A["risk_free_used"], F_PCT, BLUE, True, f"source: {A['risk_free_source']}"),
        (13, "Equity risk premium (base)", A["erp_base"], F_PCT, BLUE, True, "US implied ERP typically 4-6%"),
        (14, "Regime ERP shift", rg["erp_shift"], F_PCT, BLUE, True, f"MacroPulse regime: {rg['regime']}"),
        (15, "Equity risk premium used", "=B13+B14", F_PCT, BLACK, False, ""),
        (16, "Beta", A["beta_used"], "0.00", BLUE, True,
         f"{be['n_weeks']}w weekly regression vs S&P 500, raw {be['raw']:.2f}, Blume-adjusted"),
        (17, "Cost of equity", "=B12+B16*B15", F_PCT, BLACK, False, "CAPM"),
        (18, "Pre-tax cost of debt", A["cost_of_debt_used"], F_PCT, BLUE, True, "interest expense / debt, clipped 3-10%"),
        (19, "Tax rate", A["tax_rate"], F_PCT, BLUE, True, "21% federal + ~4% state"),
        (20, "After-tax cost of debt", "=B18*(1-B19)", F_PCT, BLACK, False, ""),
        (21, "Total debt ($m)", f"=Financials!{last_col}{R['debt']}", F_M, GREEN, False, "latest 10-K"),
        (22, "Cash ($m)", f"=Financials!{last_col}{R['cash']}", F_M, GREEN, False, ""),
        (23, "Net debt ($m)", "=B21-B22", F_M, BLACK, False, ""),
        (24, "Equity weight", "=IFERROR(B10/(B10+B21),1)", F_PCT, BLACK, False, "market values"),
        (25, "WACC", "=B24*B17+(1-B24)*B20", F_PCT, BLACK, False, ""),
        (27, "Year-1 revenue growth", g_cell, F_PCT, GREEN if A.get("growth_high") is None else BLUE, True,
         "default: historical revenue CAGR clipped to -5%..20%"),
        (28, "Terminal growth", A["terminal_growth"], F_PCT, BLUE, True, "keep at or below long-run nominal GDP"),
        (29, "FCF margin", m_cell, F_PCT, GREEN if A.get("fcf_margin") is None else BLUE, True,
         "default: 3-year average (CFO - CapEx) / revenue"),
        (30, "Projection years", N, "0", BLUE, False, "sheet layout is generated for this value"),
        (32, "Macro regime", rg["regime"], None, BLUE, False, rg["source"]),
    ]
    for r, label, v, fmt, color, key, note in rows:
        wa.cell(row=r, column=1, value=label).font = _f()
        _put(wa, f"B{r}", v, fmt, color, fill=KEY if key else None)
        wa.cell(row=r, column=3, value=note).font = _f("595959", size=9)
    wa.column_dimensions["A"].width = 28
    wa.column_dimensions["B"].width = 18
    wa.column_dimensions["C"].width = 64

    # ── DCF ──
    wd = wb.create_sheet("DCF")
    _title(wd, "DCF — FCFF proxy (CFO - CapEx), end-of-year discounting", "USD millions · growth fades linearly to terminal g")
    cols = [L(3 + i) for i in range(N)]
    lc = cols[-1]
    _hdr(wd, 4, ["Item", f"FY{years[-1]} (actual)"] + [f"FY{int(years[-1]) + i + 1}E" for i in range(N)])
    for r, t in {5: "Year t", 6: "Revenue growth", 7: "Revenue", 8: "Free cash flow", 9: "Discount factor",
                 10: "PV of FCF"}.items():
        wd.cell(row=r, column=1, value=t).font = _f()
    _put(wd, "B5", 0, "0")
    _put(wd, "B7", f"=Financials!{last_col}{rev_r}", F_M, GREEN)
    for i, c in enumerate(cols):
        p = L(2 + i)
        _put(wd, f"{c}5", i + 1, "0", BLUE)
        _put(wd, f"{c}6", f"=Assumptions!$B$27+(Assumptions!$B$28-Assumptions!$B$27)*({c}5-1)/{max(N - 1, 1)}", F_PCT)
        _put(wd, f"{c}7", f"={p}7*(1+{c}6)", F_M)
        _put(wd, f"{c}8", f"={c}7*Assumptions!$B$29", F_M)
        _put(wd, f"{c}9", f"=1/(1+Assumptions!$B$25)^{c}5", "0.0000")
        _put(wd, f"{c}10", f"={c}8*{c}9", F_M)
    out = [(12, "Sum of PV (FCF)", f"=SUM({cols[0]}10:{lc}10)", F_M),
           (13, "Terminal value (Gordon)", f"={lc}8*(1+Assumptions!B28)/(Assumptions!B25-Assumptions!B28)", F_M),
           (14, "PV of terminal value", f"=B13*{lc}9", F_M),
           (15, "Enterprise value", "=B12+B14", F_M),
           (16, "(-) Net debt", "=Assumptions!B23", F_M),
           (17, "Equity value", "=B15-B16", F_M),
           (18, "Intrinsic value per share ($)", "=B17*1E6/Assumptions!B9", F_PS),
           (19, "Share price ($)", "=Assumptions!B8", F_PS),
           (20, "Upside / (downside)", "=B18/B19-1", F_PCT),
           (21, "Terminal value % of EV", "=B14/B15", F_PCT),
           (22, "Implied EV / EBITDA", f'=IFERROR(B15/Financials!{last_col}{RR["EBITDA (EBIT + D&A)"]},"")', F_X)]
    for r, label, fml, fmt in out:
        wd.cell(row=r, column=1, value=label).font = _f(bold=r in (18, 20))
        _put(wd, f"B{r}", fml, fmt, GREEN if fml.startswith(("=Assumptions", "=Financials")) else BLACK, bold=r in (18, 20))
    wd.column_dimensions["A"].width = 30
    for c in ["B"] + cols:
        wd.column_dimensions[c].width = 14

    # ── Sensitivity ──
    ws = wb.create_sheet("Sensitivity")
    _title(ws, "Sensitivity — value per share ($): WACC (rows) x terminal growth (columns)",
           "FCF path fixed from the DCF sheet; only the discount rate and terminal growth change")
    steps = [-0.01, -0.005, 0, 0.005, 0.01]
    ws["A4"], ws["A4"].font, ws["A4"].fill = "WACC \\ g", _f(WHITE, True), HDR
    for j, s in enumerate(steps):
        c = ws.cell(row=4, column=2 + j, value=f"=Assumptions!$B$28+({s})")
        c.font, c.number_format, c.fill = _f(WHITE, True), F_PCT, HDR
    rf_, rt_ = f"DCF!${cols[0]}$8:${lc}$8", f"DCF!${cols[0]}$5:${lc}$5"
    for i, s in enumerate(steps):
        r = 5 + i
        c = ws.cell(row=r, column=1, value=f"=Assumptions!$B$25+({s})")
        c.font, c.number_format = _f(bold=True), F_PCT
        for j in range(len(steps)):
            g = f"{L(2 + j)}$4"
            cell = ws.cell(row=r, column=2 + j, value=(
                f'=IFERROR((SUMPRODUCT({rf_}/(1+$A{r})^{rt_})+DCF!${lc}$8*(1+{g})/($A{r}-{g})/(1+$A{r})^DCF!${lc}$5'
                f'-Assumptions!$B$23)*1E6/Assumptions!$B$9,"")'))
            cell.font, cell.number_format = _f(), F_PS
            if i == 2 and j == 2:
                cell.fill = KEY
    ws.cell(row=11, column=1, value="vs. share price").font = _f(bold=True)
    for i in range(len(steps)):
        c = ws.cell(row=12 + i, column=1, value=f"=A{5 + i}")
        c.font, c.number_format = _f(), F_PCT
        for j in range(len(steps)):
            cell = ws.cell(row=12 + i, column=2 + j, value=f'=IFERROR({L(2 + j)}{5 + i}/Assumptions!$B$8-1,"")')
            cell.font, cell.number_format = _f(), F_PCT
    # FOMC shock overlay (tab 13 scenarios): each one is a (WACC, g) point on the grid above.
    _hdr(ws, 18, ["FOMC scenario (tab 13)", "d rf", "d ERP", "d spread", "d g", "WACC", "g", "Value / share ($)",
                  "vs. price", "Tab-13 price shock"])
    for i, x in enumerate(val.get("fomc_overlay", [])):
        r = 19 + i
        ws.cell(row=r, column=1, value=x["title"]).font = _f()
        for col, k in ((2, "d_rf"), (3, "d_erp"), (4, "d_spread"), (5, "d_g")):
            _put(ws, f"{L(col)}{r}", x[k], '0.00%;(0.00%);"-"', BLUE)
        _put(ws, f"F{r}", f"=Assumptions!$B$24*(Assumptions!$B$12+B{r}+Assumptions!$B$16*(Assumptions!$B$15+C{r}))"
                          f"+(1-Assumptions!$B$24)*(Assumptions!$B$18+B{r}+D{r})*(1-Assumptions!$B$19)", F_PCT)
        _put(ws, f"G{r}", f"=Assumptions!$B$28+E{r}", F_PCT)
        _put(ws, f"H{r}", (f'=IFERROR((SUMPRODUCT({rf_}/(1+F{r})^{rt_})+DCF!${lc}$8*(1+G{r})/(F{r}-G{r})'
                           f'/(1+F{r})^DCF!${lc}$5-Assumptions!$B$23)*1E6/Assumptions!$B$9,"")'), F_PS, bold=True)
        _put(ws, f"I{r}", f'=IFERROR(H{r}/Assumptions!$B$8-1,"")', F_PCT)
        _put(ws, f"J{r}", x["tab13_price_shock"], F_PCT, BLUE)
    ws.cell(row=24, column=1, value=FOMC_NOTE).font = _f("595959", size=9)
    ws.column_dimensions["A"].width = 38
    for j in range(len(steps)):
        ws.column_dimensions[L(2 + j)].width = 14
    for c in "HIJ":
        ws.column_dimensions[c].width = 16

    # ── Scenarios (regime-weighted) ──
    wsc = wb.create_sheet("Scenarios")
    _title(wsc, f"Scenarios — probabilities set by the MacroPulse regime ({rg['regime']})",
           "Edit the blue deltas / probabilities; the expected value is probability-weighted")
    _hdr(wsc, 4, ["Scenario", "Growth delta", "FCF margin x", "WACC delta", "Probability", "Year-1 growth",
                  "FCF margin", "WACC", "Value / share ($)", "vs. price"])
    blk = 10
    for i, (nm, s) in enumerate(val["scenarios"].items()):
        r = 5 + i
        wsc.cell(row=r, column=1, value=nm).font = _f(bold=True)
        for col, v, fmt in ((2, s["d_growth"], F_PCT), (3, s["margin_mult"], "0.00"), (4, s["d_wacc"], F_PCT),
                            (5, s["prob"], F_PCT)):
            _put(wsc, f"{L(col)}{r}", v, fmt, BLUE, note=f"{rg['regime']} weight" if col == 5 else None)
        _put(wsc, f"F{r}", f"=Assumptions!$B$27+B{r}", F_PCT)
        _put(wsc, f"G{r}", f"=Assumptions!$B$29*C{r}", F_PCT)
        _put(wsc, f"H{r}", f"=Assumptions!$B$25+D{r}", F_PCT)
        b0 = blk + i * 6
        wsc.cell(row=b0, column=1, value=f"[{nm}] t / growth / revenue / FCF / PV").font = _f("595959", size=9)
        _put(wsc, f"B{b0 + 2}", "=DCF!B7", F_M, GREEN)
        for k in range(N):
            c, p = L(3 + k), L(2 + k)
            _put(wsc, f"{c}{b0}", k + 1, "0")
            _put(wsc, f"{c}{b0 + 1}", f"=$F${r}+(Assumptions!$B$28-$F${r})*({c}{b0}-1)/{max(N - 1, 1)}", F_PCT)
            _put(wsc, f"{c}{b0 + 2}", f"={p}{b0 + 2}*(1+{c}{b0 + 1})", F_M)
            _put(wsc, f"{c}{b0 + 3}", f"={c}{b0 + 2}*$G${r}", F_M)
            _put(wsc, f"{c}{b0 + 4}", f"={c}{b0 + 3}/(1+$H${r})^{c}{b0}", F_M)
        lk = L(2 + N)
        ev = (f"SUM(C{b0 + 4}:{lk}{b0 + 4})+{lk}{b0 + 3}*(1+Assumptions!$B$28)/($H${r}-Assumptions!$B$28)"
              f"/(1+$H${r})^{lk}{b0}")
        _put(wsc, f"I{r}", f"=({ev}-Assumptions!$B$23)*1E6/Assumptions!$B$9", F_PS, bold=True)
        _put(wsc, f"J{r}", f"=I{r}/Assumptions!$B$8-1", F_PCT)
    wsc.cell(row=8, column=1, value="Probability-weighted").font = _f(bold=True)
    _put(wsc, "E8", "=SUM(E5:E7)", F_PCT, note="must be 100%")
    _put(wsc, "I8", "=SUMPRODUCT(E5:E7,I5:I7)", F_PS, bold=True, fill=KEY)
    _put(wsc, "J8", "=I8/Assumptions!$B$8-1", F_PCT, bold=True)
    for c, w in zip("ABCDEFGHIJ", (26, 12, 12, 11, 11, 13, 11, 10, 15, 11)):
        wsc.column_dimensions[c].width = w

    # ── Technicals / Patterns / Prices ──
    wt = wb.create_sheet("Technicals")
    _title(wt, "Multi-timeframe technical snapshot",
           "MACD(12,26,9) · slow stochastic KD(14,3,3) · RSI(14) · Bollinger(20,2) · ADX(14) · CMF(20)")
    _hdr(wt, 4, ["Indicator", "Monthly", "Weekly", "Daily"])
    tech = ctx["technical"]

    def cross(x):
        return f"{x['type']} ({x['date']})" if x else "none in last 10 bars"

    trows = [("Date", lambda s: s["date"]), ("Close", lambda s: s["close"]), ("MA alignment", lambda s: s["ma_alignment"]),
             ("MACD", lambda s: f"{s['macd']['state']}, histogram {s['macd']['hist_trend']}"),
             ("MACD last cross", lambda s: cross(s["macd"]["last_cross"])),
             ("KD %K / %D", lambda s: f"{s['kd']['k'] or 0:.0f} / {s['kd']['d'] or 0:.0f} ({s['kd']['zone']})"),
             ("KD last cross", lambda s: cross(s["kd"]["last_cross"])),
             ("RSI(14)", lambda s: f"{s['rsi']['value'] or 0:.0f} ({s['rsi']['zone']})"),
             ("Bollinger %B", lambda s: s["bollinger"]["pct_b"]),
             ("Bollinger squeeze", lambda s: "yes" if s["bollinger"]["squeeze"] else "no"),
             ("ADX", lambda s: f"{s['adx']['value'] or 0:.0f} ({s['adx']['trend_strength']})"),
             ("Chaikin money flow (20)", lambda s: s["cmf_20"])]
    for i, (label, fn) in enumerate(trows):
        r = 5 + i
        wt.cell(row=r, column=1, value=label).font = _f()
        for j, tf in enumerate(("Monthly", "Weekly", "Daily")):
            c = wt.cell(row=r, column=2 + j, value=fn(tech[tf]))
            c.font = _f(BLUE)
            if label in ("Close",):
                c.number_format = F_PS
            elif label in ("Bollinger %B", "Chaikin money flow (20)"):
                c.number_format = "0.00"
    r = 5 + len(trows) + 1
    _hdr(wt, r, ["Support / resistance (daily)", "Price ($)", "Touches", ""])
    for k, lab in (("resistance", "Resistance"), ("support", "Support")):
        for j, lv in enumerate(rep.sr.get(k, []), 1):
            r += 1
            wt.cell(row=r, column=1, value=f"{lab} {j}").font = _f()
            _put(wt, f"B{r}", lv["price"], F_PS, BLUE)
            _put(wt, f"C{r}", lv["touches"], "0", BLUE)
    r += 2
    _hdr(wt, r, ["Pattern screen (daily, 20 bars)", "Value", "", ""])
    for k in ("bullish_count", "bearish_count", "bull_score", "bear_score", "overall"):
        r += 1
        wt.cell(row=r, column=1, value=k).font = _f()
        _put(wt, f"B{r}", rep.screen[k], None, BLUE)
    for c, w in zip("ABCD", (30, 34, 34, 34)):
        wt.column_dimensions[c].width = w

    wp = wb.create_sheet("Patterns")
    _title(wp, "Candlestick patterns & indicator events (daily, last 120 bars)")
    _hdr(wp, 4, ["Date", "Signal", "Direction", "Strength", "Close ($)"])
    sig = pd.concat([rep.screen["_all_candles"], rep.screen["_all_signals"]], ignore_index=True)
    sig = sig[sig["date"] >= b["prices"].index[-min(120, len(b["prices"]))]].sort_values("date", ascending=False)
    for i, rr in enumerate(sig.itertuples(), start=5):
        for j, v in enumerate([str(pd.Timestamp(rr.date).date()), rr.name, rr.direction, int(rr.strength), float(rr.price)]):
            c = wp.cell(row=i, column=1 + j, value=v)
            c.font = _f(BLUE)
            if j == 4:
                c.number_format = F_PS
    for c, w in zip("ABCDE", (12, 40, 10, 9, 12)):
        wp.column_dimensions[c].width = w

    wpx = wb.create_sheet("Prices")
    d = rep.mtf["Daily"][0].iloc[-260:]
    cols_px = [c for c in ["open", "high", "low", "close", "volume", "ma5", "ma20", "ma50", "ma200", "macd", "signal",
                           "hist", "k", "d", "rsi", "bb_up", "bb_lo", "adx", "cmf"] if c in d.columns]
    _hdr(wpx, 1, ["date"] + cols_px)
    for i, (dt, row) in enumerate(d[cols_px].iterrows(), start=2):
        wpx.cell(row=i, column=1, value=dt.date()).number_format = "yyyy-mm-dd"
        for j, c in enumerate(cols_px):
            cell = wpx.cell(row=i, column=2 + j, value=_num(row[c]))
            cell.font = _f(BLUE, size=9)
            cell.number_format = "#,##0" if c == "volume" else "#,##0.00"
    wpx.freeze_panes = "B2"

    # ── Summary (first sheet) ──
    wsu = wb.create_sheet("Summary", 0)
    syn = "SYNTHETIC DEMO DATA — not a real company · " if b.get("is_synthetic") else ""
    _title(wsu, f"{b['name']} ({b['ticker']}) — equity report summary", f"{syn}Sources: {src}")
    if syn:
        wsu["A2"].font = _f("C00000", True)
    _hdr(wsu, 4, ["Item", "Value", "Note"])
    rb, v = ctx["rule_based"], ctx["valuation"]
    summ = [("Share price ($)", "=Assumptions!B8", F_PS, GREEN, b["as_of"]),
            ("DCF value / share, base ($)", "=DCF!B18", F_PS, GREEN, "FCFF proxy, 5y + Gordon terminal"),
            ("Upside, base", "=DCF!B20", F_PCT, GREEN, ""),
            ("Probability-weighted value ($)", "=Scenarios!I8", F_PS, GREEN, "bear/base/bull weighted by regime"),
            ("Upside, probability-weighted", "=Scenarios!J8", F_PCT, GREEN, ""),
            ("WACC", "=Assumptions!B25", F_PCT, GREEN, ""),
            ("Terminal growth", "=Assumptions!B28", F_PCT, GREEN, ""),
            ("Sensitivity range ($)", '=TEXT(MIN(Sensitivity!B5:F9),"$#,##0.00")&" - "&TEXT(MAX(Sensitivity!B5:F9),"$#,##0.00")',
             None, GREEN, "WACC ±1pp x g ±1pp"),
            ("Macro regime", rg["regime"], None, BLUE,
             f"weights bear/base/bull {rg['probs']['Bear']:.0%}/{rg['probs']['Base']:.0%}/{rg['probs']['Bull']:.0%}, "
             f"ERP shift {rg['erp_shift']:+.2%}"),
            ("52-week range ($)", f"{v['range_52w']['low']:,.2f} - {v['range_52w']['high']:,.2f}", None, BLUE, "")]
    for mname, bd in v["relative_bands"].items():
        summ.append((f"{mname} band value 10/50/90 ($)", f"{bd['low']:,.2f} / {bd['mid']:,.2f} / {bd['high']:,.2f}",
                     None, BLUE, "historical multiple percentiles x current fundamentals"))
    for tf in ("Monthly", "Weekly", "Daily"):
        s = ctx["technical"][tf]
        summ.append((f"{tf} trend", f"{s['ma_alignment']}, MACD {s['macd']['state']}", None, BLUE, ""))
    summ += [("Pattern screen (daily)", f"{rep.screen['overall']} — bull {rep.screen['bull_score']} : bear "
              f"{rep.screen['bear_score']}", None, BLUE, "candles + indicator events, strength-weighted"),
             ("Rule-based rating", rb["rating"], None, BLUE, rb["basis"])]
    for i, (label, vv, fmt, color, note) in enumerate(summ):
        r = 5 + i
        wsu.cell(row=r, column=1, value=label).font = _f(bold=label.startswith(("DCF value", "Probability", "Rule")))
        _put(wsu, f"B{r}", vv, fmt, color)
        wsu.cell(row=r, column=3, value=note).font = _f("595959", size=9)
        if i % 2:
            for c in "ABC":
                wsu[f"{c}{r}"].fill = STRIPE
    nr = 5 + len(summ) + 1
    for k, t in enumerate(v["notes"] + ["Not investment advice. Results depend heavily on the yellow assumptions."]):
        wsu.cell(row=nr + k, column=1, value=f"* {t}").font = _f("C00000" if k < len(v["notes"]) else "595959", size=9)
    wsu.column_dimensions["A"].width = 32
    wsu.column_dimensions["B"].width = 34
    wsu.column_dimensions["C"].width = 52
    if chart_png:
        img = XLImage(io.BytesIO(chart_png))
        img.width, img.height = img.width * 0.5, img.height * 0.5
        wsu.add_image(img, "E4")
    for s in wb.worksheets:
        s.sheet_view.showGridLines = False
    wb.calculation.fullCalcOnLoad = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
