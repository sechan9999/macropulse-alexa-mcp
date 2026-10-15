"""
tests/test_equity_valuation.py
─────────────────────────────────────────────────────────────────
DCF fixes from a review of an AVGO report (fair value $67 vs a $338 price):
  * base = trailing twelve months from the latest 10-Q, net debt at that quarter end
  * discounting from the price date (stub), FCFF with after-tax interest added back
  * year-1 growth from the 10-Q year-to-date, decaying with a half-life instead of a 5-year linear cliff
  * reverse DCF, scenario growth deltas scaled to the growth rate, valuation half weight when the
    bear-bull range does not contain the price
  * split-adjusted EPS for the P/E band, blended weekly/monthly beta
Offline: SEC facts, prices and splits are injected.
"""
import copy
import io
import shutil
import subprocess
import tempfile
import unittest
from dataclasses import asdict
from pathlib import Path

import numpy as np
import pandas as pd

from src.equity_report.config import DCFAssumptions
from src.equity_report.engine import analysis, data_us, dcf, synthetic

SPLIT_10_FOR_1 = pd.Series([10.0], index=[pd.Timestamp("2025-06-01")])


def _fixture_bundle(splits=SPLIT_10_FOR_1, facts=None):
    facts = facts or synthetic.company_facts_fixture()

    def http_get(url):
        if url.endswith("company_tickers.json"):
            return {"0": {"cik_str": 1234567, "ticker": "FIX", "title": "FIXTURE INC"}}
        return facts
    data_us._mem.clear()
    return data_us.load_bundle("FIX", price_fetch=lambda t: synthetic.price_frame(seed=3 if t == "FIX" else 5),
                               http_get=http_get, with_positioning=False, splits_fetch=lambda t: splits)


class TestTrailingTwelveMonths(unittest.TestCase):
    def setUp(self):
        self.facts = synthetic.company_facts_fixture()
        self.fin = data_us.annual_table(self.facts)

    def test_ttm_is_fiscal_year_plus_ytd_minus_prior_ytd(self):
        t = data_us.ttm(self.facts, self.fin)
        self.assertEqual(t["period_end"], "2026-03-28")
        self.assertEqual(t["ytd_months"], 6)
        self.assertEqual(t["revenue"], 130 + 70 - 60)                 # the 3-month row (37) is not used
        self.assertEqual(t["operating_cf"], 38 + 22 - 18)
        self.assertEqual(t["capex"], 9 + 5 - 4)
        self.assertAlmostEqual(t["interest_expense"], 2.5 + 1.4 - 1.2)
        self.assertAlmostEqual(t["ytd_growth"], 70 / 60 - 1)
        self.assertEqual((t["cash"], t["debt"]), (30.0, 52.0))       # quarter-end balance sheet, 48 + 4

    def test_interest_missing_from_the_10q_falls_back_to_the_fiscal_year(self):
        facts = copy.deepcopy(self.facts)
        rows = facts["facts"]["us-gaap"]["InterestExpense"]["units"]["USD"]
        facts["facts"]["us-gaap"]["InterestExpense"]["units"]["USD"] = [r for r in rows if r["form"] != "10-Q"]
        t = data_us.ttm(facts, data_us.annual_table(facts))
        self.assertEqual(t["interest_expense"], 2.5)                  # FY2025, not 0
        self.assertTrue(t["interest_expense_from_fy"])
        self.assertFalse(data_us.ttm(self.facts, self.fin)["interest_expense_from_fy"])

    def test_no_10q_after_the_10k_means_no_ttm(self):
        facts = copy.deepcopy(self.facts)
        for tag in facts["facts"]["us-gaap"].values():
            tag["units"] = {u: [r for r in rows if r.get("form") != "10-Q"] for u, rows in tag["units"].items()}
        self.assertIsNone(data_us.ttm(facts, data_us.annual_table(facts)))


class TestSplitAdjustedEps(unittest.TestCase):
    def test_eps_filed_before_a_split_is_divided_by_it(self):
        eps = data_us.eps_by_filing_date(synthetic.company_facts_fixture(), SPLIT_10_FOR_1)
        self.assertEqual(list(eps.round(4).values), [0.27, 3.00])    # 2.70 filed 2024-11-01, before the split

    def test_annual_table_eps_is_split_adjusted(self):
        facts = synthetic.company_facts_fixture()
        fin = data_us.split_adjust_eps(data_us.annual_table(facts), facts, SPLIT_10_FOR_1)
        self.assertAlmostEqual(fin.loc["2025", "eps_diluted"], 3.00)  # filed after the split: unchanged

    def test_bundle_uses_ttm_eps(self):
        b = _fixture_bundle()
        self.assertAlmostEqual(b["eps_ttm"], (30 + 17 - 14) / (1000.0 * 1010 / 990))
        self.assertIn("10-Q", b["sources"][-1])


class TestDilution(unittest.TestCase):
    def test_latest_quarter_ratio(self):
        f, period = data_us.dilution_factor(synthetic.company_facts_fixture())
        self.assertAlmostEqual(f, 1010 / 990)                       # 3-month Q2 FY26, not the 6-month YTD
        self.assertEqual(period, "2025-12-28 to 2026-03-28")

    def test_missing_counts_mean_no_dilution(self):
        self.assertEqual(data_us.dilution_factor({"facts": {"us-gaap": {}}}), (1.0, None))

    def test_per_share_values_use_diluted_shares_and_market_cap_basic(self):
        b = _fixture_bundle()
        v = dcf.build(b)
        self.assertAlmostEqual(b["shares_diluted"], 1000 * 1010 / 990)
        self.assertAlmostEqual(v["inputs"]["shares"], b["shares_diluted"])
        self.assertAlmostEqual(v["base"]["per_share"], v["base"]["equity"] / b["shares_diluted"])
        self.assertAlmostEqual(v["market_cap"], b["last_price"] * 1000)
        self.assertEqual(dcf.build(synthetic.stock_bundle())["assumptions"]["dilution_factor"], 1.0)


class TestDcfMechanics(unittest.TestCase):
    def test_stub_discounting_scales_ev(self):
        inp = dcf.DCFInputs(1000, 0.2, 0.1, 0.02, 0.09, 10, 0, 10, 100, 0.0, 2.0)
        ev0 = dcf.value(inp)["ev"]
        ev = dcf.value(dcf.DCFInputs(**{**asdict(inp), "stub": 0.4}))["ev"]
        self.assertAlmostEqual(ev, ev0 * 1.09 ** 0.4)

    def test_growth_decay_path(self):
        g = dcf.growth_path(0.5, 0.025, 10, 2.0)
        self.assertAlmostEqual(g[0], 0.5)
        self.assertAlmostEqual(g[-1], 0.025)
        self.assertTrue((np.diff(g) < 0).all())
        excess = g - 0.025
        self.assertAlmostEqual(excess[2] / excess[0], 0.5, delta=0.03)  # halves over the 2-year half-life
        lin = dcf.growth_path(0.5, 0.025, 5)
        self.assertAlmostEqual(lin[1] - lin[0], (0.025 - 0.5) / 4)       # half_life None -> linear (old behaviour)

    def test_fcff_adds_back_after_tax_interest(self):
        b = synthetic.stock_bundle()
        v = dcf.build(b)
        last = b["financials"].iloc[-1]
        self.assertAlmostEqual(v["history"]["fcf"].iloc[-1],
                               last["operating_cf"] - last["capex"] + last["interest_expense"] * 0.75)

    def test_reverse_dcf_reproduces_the_price(self):
        inp = dcf.DCFInputs(1000, 0.25, 0.10, 0.025, 0.10, 10, 200, 10, 450, 0.3, 2.0)
        g = dcf.implied_growth(inp)
        w = dcf.implied_wacc(inp)
        self.assertAlmostEqual(dcf.value(dcf.DCFInputs(**{**asdict(inp), "growth_high": g}))["per_share"], 450, places=6)
        self.assertAlmostEqual(dcf.value(dcf.DCFInputs(**{**asdict(inp), "wacc": w}))["per_share"], 450, places=6)

    def test_scenario_growth_delta_scales(self):
        self.assertEqual(dcf.scenario_growth_delta(0.08), 0.05)
        self.assertAlmostEqual(dcf.scenario_growth_delta(0.50), 0.15)


class TestBuildOnTtm(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.b = _fixture_bundle()
        cls.v = dcf.build(cls.b)

    def test_base_period_inputs(self):
        a, inp = self.v["assumptions"], self.v["inputs"]
        self.assertEqual(a["base_period"], "TTM to 2026-03-28")
        self.assertEqual(inp["revenue0"], 140.0)
        self.assertAlmostEqual(inp["net_debt"], 52.0 - 30.0)
        self.assertAlmostEqual(a["growth_high_used"], 70 / 60 - 1)
        self.assertAlmostEqual(a["fcf_margin_used"], (42 - 10 + 2.7 * 0.75) / 140)
        self.assertAlmostEqual(inp["stub"], (pd.Timestamp(self.b["as_of"]) - pd.Timestamp("2026-03-28")).days / 365.25)
        self.assertEqual(inp["years"], 10)

    def test_overrides_win(self):
        v = dcf.build(self.b, DCFAssumptions(growth_high=0.3, fcf_margin=0.4, beta=1.2))
        a = v["assumptions"]
        self.assertEqual((a["growth_high_used"], a["fcf_margin_used"], a["beta_used"]), (0.3, 0.4, 1.2))

    def test_reverse_dcf_note(self):
        self.assertTrue(any(n.startswith("Reverse DCF") for n in self.v["notes"]))


class TestRating(unittest.TestCase):
    def test_valuation_has_half_weight_when_the_scenarios_miss_the_price(self):
        ctx = copy.deepcopy(analysis.analyze(synthetic.stock_bundle()).ctx)
        ctx["valuation"]["expected_upside"] = -0.6
        ctx["valuation"]["scenarios_bracket_price"] = True
        full = analysis.rule_opinion(ctx)
        ctx["valuation"]["scenarios_bracket_price"] = False
        half = analysis.rule_opinion(ctx)
        self.assertAlmostEqual(half["score"] - full["score"], 1.5)
        self.assertIn("half weight", half["basis"])


class TestContinuousValuationScore(unittest.TestCase):
    def _score(self, up, bracket=True):
        ctx = copy.deepcopy(analysis.analyze(synthetic.stock_bundle()).ctx)
        ctx["valuation"]["expected_upside"], ctx["valuation"]["scenarios_bracket_price"] = up, bracket
        base = copy.deepcopy(ctx)
        base["valuation"]["expected_upside"] = 0.0
        return analysis.rule_opinion(ctx)["score"] - analysis.rule_opinion(base)["score"]

    def test_one_point_per_ten_percent_capped_at_three(self):
        self.assertAlmostEqual(self._score(-0.12), -1.2)
        self.assertAlmostEqual(self._score(0.05), 0.5)
        self.assertAlmostEqual(self._score(-0.60), -3.0)
        self.assertAlmostEqual(self._score(0.40), 3.0)

    def test_no_cliff_at_ten_percent(self):
        self.assertLess(abs(self._score(-0.101) - self._score(-0.099)), 0.05)

    def test_half_weight_still_applies(self):
        self.assertAlmostEqual(self._score(-0.12, bracket=False), -0.6)


class TestBeta(unittest.TestCase):
    def test_blend_of_weekly_and_monthly(self):
        idx = synthetic.price_frame(n=1300, start=4000, drift=0.0004, vol=0.010, seed=5)
        mkt = idx["Close"].pct_change().fillna(0).to_numpy()
        stock = synthetic.price_frame(n=1300, market=mkt, beta=1.4, seed=8)
        b = dcf.estimate_beta(stock["Close"], idx["Close"])
        self.assertGreaterEqual(b["n_months"], 36)
        self.assertAlmostEqual(b["raw"], (b["raw_weekly"] + b["raw_monthly"]) / 2)
        self.assertAlmostEqual(b["adjusted"], 0.67 * b["raw"] + 0.33)

    def test_short_history_uses_weekly_only(self):
        d = synthetic.stock_bundle()
        b = dcf.estimate_beta(d["prices"]["close"], d["index"]["close"])
        self.assertIsNone(b["raw_monthly"])
        self.assertAlmostEqual(b["raw"], b["raw_weekly"])


@unittest.skipIf(shutil.which("soffice") is None, "LibreOffice not installed")
class TestExcelOnTtm(unittest.TestCase):
    def test_libreoffice_recalc_matches_python(self):
        from openpyxl import load_workbook
        from src.equity_report.render import excel
        rep = analysis.analyze(_fixture_bundle())
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "m.xlsx"
            src.write_bytes(excel.build(rep))
            subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(Path(d) / "calc"), str(src)],
                           check=True, capture_output=True, timeout=180)
            wb = load_workbook(Path(d) / "calc" / "m.xlsx", data_only=True)
        v = rep.val
        self.assertAlmostEqual(wb["Assumptions"]["B25"].value, v["assumptions"]["wacc"], places=9)
        self.assertAlmostEqual(wb["DCF"]["B7"].value * 1e6, v["inputs"]["revenue0"], places=6)
        self.assertAlmostEqual(wb["DCF"]["C6"].value, v["base"]["growth"][0], places=9)
        self.assertAlmostEqual(wb["DCF"]["L6"].value, v["base"]["growth"][-1], places=9)
        self.assertAlmostEqual(wb["DCF"]["B18"].value, v["base"]["per_share"], places=6)
        self.assertAlmostEqual(wb["Scenarios"]["I8"].value, v["expected_per_share"], places=6)
        sens = [wb["Sensitivity"].cell(r, c).value for r in range(5, 10) for c in range(2, 7)]
        self.assertAlmostEqual(min(sens), v["sensitivity"].values.min(), places=6)
        for i, x in enumerate(v["fomc_overlay"]):
            self.assertAlmostEqual(wb["Sensitivity"][f"H{19 + i}"].value, x["per_share"], places=6)
        self.assertEqual(wb["Financials"]["A27"].value, "FCFF (CFO - CapEx + interest x (1 - tax))")
        self.assertAlmostEqual(wb["Assumptions"]["B40"].value, v["inputs"]["shares"], places=6)


if __name__ == "__main__":
    unittest.main()
