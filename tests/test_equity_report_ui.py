"""
tests/test_equity_report_ui.py
─────────────────────────────────────────────────────────────────
Renderers + Streamlit tab, offline (DEMO ticker = synthetic data):
  * build_all() returns Excel / Word / HTML / zip bytes with the expected contents
  * the Excel model, recalculated by LibreOffice, reproduces the Python DCF (skipped without soffice)
  * the Equity Report tab renders, shows KPIs and four download buttons
  * regime_inputs() refuses stale or demo macro frames
"""
import io
import shutil
import subprocess
import tempfile
import unittest
import zipfile
from datetime import datetime, timedelta
from pathlib import Path

import pandas as pd

from src.equity_report.engine import pipeline
from src.equity_report.render import build_all
from src.equity_report.ui import regime_inputs


class TestRenderers(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.rep = pipeline.run("DEMO", regime_label="Risk-Off 🔴")
        cls.files = build_all(cls.rep, dashboard_theme="dark")

    def test_all_files_present(self):
        f = self.files
        self.assertTrue(f["xlsx"][:2] == b"PK" and f["docx"][:2] == b"PK" and f["zip"][:2] == b"PK")
        self.assertIn("SYNTHETIC DEMO DATA", f["html"])
        self.assertIn('data-theme="dark"', f["html_embed"])
        self.assertIn("Risk-Off", f["html"])
        self.assertIn("FOMC shock overlay", f["html"])
        self.assertIn('class="fomc"', f["html"])
        names = zipfile.ZipFile(io.BytesIO(f["zip"])).namelist()
        for n in (f["filenames"]["xlsx"], f["filenames"]["docx"], f["filenames"]["html"], "summary.json",
                  "charts/DEMO_daily.png"):
            self.assertIn(n, names)
        self.assertTrue(all(n.isascii() for n in names))

    def test_word_has_sections(self):
        from docx import Document
        text = "\n".join(p.text for p in Document(io.BytesIO(self.files["docx"])).paragraphs)
        for h in ("1. Fundamentals", "3. Technical Analysis", "5. Macro Regime Overlay", "6. Verdict"):
            self.assertIn(h, text)

    @unittest.skipIf(shutil.which("soffice") is None, "LibreOffice not installed")
    def test_excel_recalc_matches_python(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as d:
            src = Path(d) / "m.xlsx"
            src.write_bytes(self.files["xlsx"])
            subprocess.run(["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(Path(d) / "calc"), str(src)],
                           check=True, capture_output=True, timeout=180)
            wb = load_workbook(Path(d) / "calc" / "m.xlsx", data_only=True)
        v = self.rep.val
        self.assertAlmostEqual(wb["Assumptions"]["B25"].value, v["assumptions"]["wacc"], places=9)
        self.assertAlmostEqual(wb["DCF"]["B18"].value, v["base"]["per_share"], places=6)
        self.assertAlmostEqual(wb["Scenarios"]["I8"].value, v["expected_per_share"], places=6)
        sens = [wb["Sensitivity"].cell(r, c).value for r in range(5, 10) for c in range(2, 7)]
        self.assertAlmostEqual(min(sens), v["sensitivity"].values.min(), places=6)
        self.assertAlmostEqual(max(sens), v["sensitivity"].values.max(), places=6)
        for i, x in enumerate(v["fomc_overlay"]):                     # FOMC overlay block, rows 19-22
            self.assertAlmostEqual(wb["Sensitivity"][f"F{19 + i}"].value, x["wacc"], places=9)
            self.assertAlmostEqual(wb["Sensitivity"][f"H{19 + i}"].value, x["per_share"], places=6)

    def test_excel_fomc_block(self):
        from openpyxl import load_workbook
        ws = load_workbook(io.BytesIO(self.files["xlsx"]))["Sensitivity"]
        self.assertEqual(ws["A19"].value, "Hawkish Surprise (+50 bps Hike)")
        self.assertTrue(str(ws["H19"].value).startswith("=IFERROR((SUMPRODUCT("))
        self.assertTrue(str(ws["F22"].value).startswith("=Assumptions!$B$24"))

    def test_word_has_fomc_table(self):
        from docx import Document
        cells = {c.text for t in Document(io.BytesIO(self.files["docx"])).tables for r in t.rows for c in r.cells}
        self.assertIn("FOMC scenario (tab 13)", cells)


class TestRegimeInputs(unittest.TestCase):
    def _frame(self, end, demo=False):
        idx = pd.date_range(end=end, periods=3, freq="MS")
        return pd.DataFrame({"regime": ["Neutral 🟡", "Neutral 🟡", "Risk-Off 🔴"], "dgs10": [4.1, 4.2, 4.35],
                             "_is_demo": demo}, index=idx)

    def test_live_frame(self):
        reg, rf, _ = regime_inputs(self._frame(datetime.now()))
        self.assertEqual(reg, "Risk-Off 🔴")
        self.assertAlmostEqual(rf, 0.0435)

    def test_stale_frame_goes_neutral(self):
        reg, rf, note = regime_inputs(self._frame(datetime.now() - timedelta(days=400)))
        self.assertIsNone(reg)
        self.assertIsNone(rf)
        self.assertIn("days old", note)

    def test_demo_frame_goes_neutral(self):
        reg, rf, _ = regime_inputs(self._frame(datetime.now(), demo=True))
        self.assertIsNone(reg)


def _deeplink_script():
    import streamlit as st
    import streamlit.components.v1 as components
    from src.equity_report.ui import render_tab
    st.session_state.setdefault("er_provider", "Rules only (no LLM)")
    render_tab(st, components, None)


def _tab_script():
    import streamlit as st
    import streamlit.components.v1 as components
    from src.equity_report.ui import render_tab
    render_tab(st, components, None)


class TestStreamlitTab(unittest.TestCase):
    def test_tab_renders_demo_with_downloads(self):
        from streamlit.testing.v1 import AppTest
        at = AppTest.from_function(_tab_script, default_timeout=120)
        at.run()
        self.assertFalse(at.exception)
        at.text_input(key="er_ticker").set_value("DEMO")
        at.selectbox(key="er_provider").set_value("Rules only (no LLM)")
        at.button(key="er_go").click()
        at.run()
        self.assertFalse(at.exception, at.exception)
        self.assertEqual(len(at.metric), 6)
        self.assertTrue(any("SYNTHETIC" in w.value for w in at.warning))
        downloads = at.get("download_button")
        self.assertEqual(len(downloads), 4)

    def test_ticker_deep_link_generates_without_click(self):
        from streamlit.testing.v1 import AppTest
        at = AppTest.from_function(_deeplink_script, default_timeout=120)
        at.query_params["ticker"] = "demo"
        at.run()
        self.assertFalse(at.exception, at.exception)
        self.assertEqual(len(at.metric), 6)                           # report built on arrival
        self.assertEqual(len(at.get("download_button")), 4)
        at.run()                                                      # a rerun does not rebuild / re-trigger
        self.assertEqual(at.session_state["er_autorun_done"], "DEMO")


if __name__ == "__main__":
    unittest.main()
